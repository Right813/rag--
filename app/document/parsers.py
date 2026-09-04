import logging
import re
from pathlib import Path

from app.document.models import ParsedPage

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".markdown", ".jpg", ".jpeg", ".png"}


class DocumentParseError(ValueError):
    pass


def _ocr_required_message() -> str:
    return "该文件需要 OCR 解析，请设置 OCR_ENABLED=true 并安装 requirements-optional.txt"


class DocumentParser:
    def __init__(self, ocr_enabled: bool = True) -> None:
        self.ocr_enabled = ocr_enabled

    def parse(self, file_path: str | Path) -> list[ParsedPage]:
        path = Path(file_path)
        suffix = path.suffix.casefold()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise DocumentParseError(f"不支持的文件类型：{suffix or 'unknown'}")
        if suffix == ".pdf":
            return self._parse_pdf(path)
        if suffix == ".docx":
            return self._parse_docx(path)
        if suffix == ".xlsx":
            return self._parse_xlsx(path)
        if suffix in {".txt", ".md", ".markdown"}:
            return self._parse_text(path)
        return self._parse_image(path)

    @staticmethod
    def _parse_text(path: Path) -> list[ParsedPage]:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="gb18030")
        return [ParsedPage(content=text, page=1)]

    def _parse_pdf(self, path: Path) -> list[ParsedPage]:
        try:
            import fitz
        except ImportError as exc:
            raise DocumentParseError("PDF 解析依赖 PyMuPDF 未安装") from exc
        pages: list[ParsedPage] = []
        with fitz.open(path) as document:
            for page_number, page in enumerate(document, start=1):
                text = page.get_text("text") or ""
                if text.strip():
                    pages.append(ParsedPage(content=text, page=page_number))
                    continue
                ocr_text = self._ocr_image(page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))) if self.ocr_enabled else ""
                pages.append(ParsedPage(content=ocr_text, page=page_number))
        if not any(page.content.strip() for page in pages):
            if not self.ocr_enabled:
                raise DocumentParseError(_ocr_required_message())
            raise DocumentParseError("PDF 未解析出可用文本；请安装 PaddleOCR 处理扫描件")
        return pages

    @staticmethod
    def _parse_docx(path: Path) -> list[ParsedPage]:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentParseError("DOCX 解析依赖 python-docx 未安装") from exc
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            if rows:
                paragraphs.append(f"表格 {table_index}\n" + "\n".join(rows))
        return [ParsedPage(content="\n".join(paragraphs), page=1)]

    @staticmethod
    def _parse_xlsx(path: Path) -> list[ParsedPage]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise DocumentParseError("XLSX 解析依赖 openpyxl 未安装") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        pages: list[ParsedPage] = []
        try:
            for sheet in workbook.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        rows.append(" | ".join(values))
                if rows:
                    pages.append(ParsedPage(content=f"工作表：{sheet.title}\n" + "\n".join(rows), page=len(pages) + 1))
        finally:
            workbook.close()
        if not pages:
            raise DocumentParseError("XLSX 没有可用单元格内容")
        return pages

    def _parse_image(self, path: Path) -> list[ParsedPage]:
        if not self.ocr_enabled:
            raise DocumentParseError(_ocr_required_message())
        try:
            from paddleocr import PaddleOCR
        except ImportError as exc:
            raise DocumentParseError("图片解析依赖 PaddleOCR 未安装") from exc
        try:
            ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
            result = ocr.ocr(str(path), cls=True)
            lines: list[str] = []
            for page in result or []:
                for item in page or []:
                    if len(item) > 1 and item[1]:
                        lines.append(str(item[1][0]))
            text = "\n".join(lines)
        except Exception as exc:
            logger.exception("OCR failed for %s", path)
            raise DocumentParseError(f"OCR 解析失败：{exc}") from exc
        if not text.strip():
            raise DocumentParseError("图片未识别出可用文本")
        return [ParsedPage(content=text, page=1)]

    @staticmethod
    def _ocr_image(pixmap) -> str:
        try:
            from paddleocr import PaddleOCR
        except ImportError:
            return ""
        try:
            ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
            result = ocr.ocr(pixmap.tobytes("png"), cls=True)
            lines = []
            for page in result or []:
                for item in page or []:
                    if len(item) > 1 and item[1]:
                        lines.append(str(item[1][0]))
            return "\n".join(lines)
        except Exception:
            logger.exception("PDF OCR failed")
            return ""


def detect_structure(text: str, current_chapter: str = "", current_section: str = "") -> tuple[str, str]:
    value = re.sub(r"\s+", " ", text.strip())
    if re.match(r"^(第[一二三四五六七八九十百零0-9]+[章节篇]|[一二三四五六七八九十]+、)", value):
        return value, ""
    if re.match(r"^\d+(?:\.\d+){1,3}\s*[、.．)]?\s*\S+", value) or re.match(r"^[一二三四五六七八九十]+、\s*\S+", value):
        return current_chapter, value
    return current_chapter, current_section
